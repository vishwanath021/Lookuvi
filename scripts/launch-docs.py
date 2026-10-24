#!/usr/bin/env python3
"""Turn docs/test-plan/results.json into markdown, CSV, XLSX, and print-ready HTML."""
import csv
import html
import json
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

root = Path("docs/test-plan")
data = json.loads((root / "results.json").read_text())
results = data["results"]
summary = data["summary"]

by_func = OrderedDict()
for row in results:
    by_func.setdefault(row["functionality"], []).append(row)

def esc(value):
    return html.escape(str(value or "")).replace("\n", "<br>")

lines = []
lines.append("# Lookuvi launch test plan")
lines.append("")
lines.append("Executed on 24 October 2026 against the local server. AI provider is mock. WhatsApp is not a live handset: webhook signatures use a local secret, and reminder sends use an injected sender that never calls Meta. `OPENAI_API_KEY` was removed from the process. Demo logins only.")
lines.append("")
lines.append(f"Result: **{summary['pass']} pass**, **{summary['fail']} fail**, **{summary['untested']} untested**, {summary['total']} cases.")
lines.append("")
lines.append("The inventory counts from the approval note were sizing estimates. This sheet is the case list that was actually written and run. A case that was not run is UNTESTED with a reason. Nothing here is a silent pass.")
lines.append("")
lines.append("## Onboarding checklist")
lines.append("")
lines.append("Use this for each of the seven salons. Do it in order.")
lines.append("")
checklist = [
    ("1", "Super admin creates the salon", "Name, lowercase slug, city, WhatsApp, owner email, a password of at least 12 characters that is not a demo password. Plan as sold. Queue is already on. One trial grant. Styles are copied."),
    ("2", "Owner signs in and finishes onboarding", "Brand, services, WhatsApp, languages, then Finish. Staff can see the wizard until Finish."),
    ("3", "Owner sets chairs and hours", "On the queue desk, not in Settings. Chairs 1–30, average minutes, live line or time requests, chair/station/room. Pause the line before the QR is public if the shop is not ready."),
    ("4", "Owner checks caps", "Anonymous daily, member daily, monthly try-ons, visit tries for bridal. Daily and monthly reset at midnight Asia/Kolkata."),
    ("5", "Owner confirms reminders", "On by default, 35 days, range 28–42. WhatsApp monthly cap default 50. Reminders and offers count. Queue alerts and the desk WhatsApp button do not."),
    ("6", "Print the QR only on the real domain", "The code points at the welcome page. Localhost refuses to print and says it is not ready."),
    ("7", "Staff rehearsal", "Join from a phone, Next, Start, Done. In request mode: confirm, decline, Mark visited. Walk-in only on the live line."),
    ("8", "Book this look", "A guest tap creates a lead and opens WhatsApp. Lite does not download the CSV. Starter and Pro do."),
    ("9", "STOP rehearsal", "A reply of STOP opts out that salon only. STOPALL or UNSUBSCRIBE opts out the shared number everywhere. The owner cannot undo STOP."),
    ("10", "Leave billing, bookings, guest hub, custom domain, outlets, salon AI key, embed, and the guide off", "Tell the owner those are not in this launch."),
]
lines.append("| Step | Do this | What good looks like |")
lines.append("| --- | --- | --- |")
for step, title, detail in checklist:
    lines.append(f"| {step} | {title} | {detail} |")
lines.append("")
lines.append("## How a case was run")
lines.append("")
lines.append("- `http` is a request to the local Next server.")
lines.append("- `http+db` checks the row afterwards.")
lines.append("- `library` calls the same function the server uses, with a fake WhatsApp sender where a real send would cost money or need Meta.")
lines.append("- `db` reads the seeded row.")
lines.append("- `script` is `npm run scan:secrets`.")
lines.append("- `k6` is the browse script. It is UNTESTED because k6 is not installed here.")
lines.append("")
lines.append("## Integration")
lines.append("")
lines.append("Mock AI returns a JPEG and sets `x-demo-reason: no-key`. The credit balance does not move. A succeeded-cap test inserts a SUCCEEDED row, because a mock sample is stored as DEMO and does not count. The WhatsApp webhook is signed with the local app secret. Cron uses the local bearer secret. Neither path calls Meta or OpenAI.")
lines.append("")
lines.append("Performance: `scripts/k6/guest-browse.js` loads health, the welcome page, and the public queue for the demo salon. Five virtual users, 30 seconds. Thresholds: failed requests under 5 percent, 95th percentile under 1.5 seconds. Not run in this environment.")
lines.append("")
lines.append("## Results by functionality")
lines.append("")
lines.append("| Id | Pass | Fail | Untested |")
lines.append("| --- | --- | --- | --- |")
for func, rows in by_func.items():
    passed = sum(1 for row in rows if row["status"] == "PASS")
    failed = sum(1 for row in rows if row["status"] == "FAIL")
    untested = sum(1 for row in rows if row["status"] == "UNTESTED")
    lines.append(f"| {func} | {passed} | {failed} | {untested} |")
lines.append("")
lines.append("## Cases")
lines.append("")
for func, rows in by_func.items():
    lines.append(f"### {func}")
    lines.append("")
    for row in rows:
        lines.append(f"#### {row['id']} {row['title']} ({row['status']})")
        lines.append("")
        lines.append(f"- Priority: {row['priority']}. Area: {row['area']}. Method: {row['method']}.")
        lines.append(f"- Given: {row['given']}")
        lines.append(f"- When: {row['when']}")
        lines.append(f"- Then: {row['then']}")
        if row.get("detail"):
            lines.append(f"- Detail: {row['detail']}")
        lines.append("")

(root / "TEST-PLAN.md").write_text("\n".join(lines))

fields = ["id", "functionality", "priority", "area", "title", "given", "when", "then", "status", "method", "detail"]
with (root / "cases.csv").open("w", newline="") as handle:
    writer = csv.DictWriter(handle, fieldnames=fields)
    writer.writeheader()
    for row in results:
        writer.writerow({key: row.get(key, "") for key in fields})

book = Workbook()
sheet = book.active
sheet.title = "Results"
header_fill = PatternFill("solid", fgColor="1F3A34")
header_font = Font(color="FFFFFF", bold=True)
for col, key in enumerate(fields, 1):
    cell = sheet.cell(1, col, key)
    cell.fill = header_fill
    cell.font = header_font
fills = {
    "PASS": PatternFill("solid", fgColor="E5F4EA"),
    "FAIL": PatternFill("solid", fgColor="F8D7DA"),
    "UNTESTED": PatternFill("solid", fgColor="FFF3CD"),
}
for index, row in enumerate(results, 2):
    for col, key in enumerate(fields, 1):
        cell = sheet.cell(index, col, row.get(key, ""))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if key == "status":
            cell.fill = fills.get(row["status"], PatternFill())
widths = [12, 14, 10, 12, 42, 36, 36, 42, 12, 12, 42]
for col, width in enumerate(widths, 1):
    sheet.column_dimensions[get_column_letter(col)].width = width
sheet.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{len(results)+1}"
sheet.freeze_panes = "A2"

check = book.create_sheet("Onboarding")
for col, name in enumerate(["Step", "Do this", "What good looks like"], 1):
    cell = check.cell(1, col, name)
    cell.fill = header_fill
    cell.font = header_font
for index, item in enumerate(checklist, 2):
    for col, value in enumerate(item, 1):
        check.cell(index, col, value).alignment = Alignment(wrap_text=True, vertical="top")
check.column_dimensions["A"].width = 8
check.column_dimensions["B"].width = 42
check.column_dimensions["C"].width = 88

book.save(root / "cases.xlsx")

body = []
body.append("<!doctype html><html><head><meta charset='utf-8'><title>Lookuvi launch test plan</title>")
body.append("<style>body{font-family:Georgia,serif;max-width:900px;margin:32px auto;color:#1c1c1c;line-height:1.45}table{border-collapse:collapse;width:100%;font-size:12px}td,th{border:1px solid #ccc;padding:6px;vertical-align:top}th{background:#1F3A34;color:white}h1,h2,h3{font-family:Palatino,serif}.pass{color:#146c43}.fail{color:#9b1c1c}.untested{color:#8a5a00}code{font-family:ui-monospace,monospace}</style></head><body>")
body.append(f"<h1>Lookuvi launch test plan</h1><p>{esc(lines[2])}</p>")
body.append(f"<p><strong>{summary['pass']} pass</strong>, {summary['fail']} fail, {summary['untested']} untested, {summary['total']} cases.</p>")
body.append("<h2>Onboarding checklist</h2><table><tr><th>Step</th><th>Do this</th><th>What good looks like</th></tr>")
for step, title, detail in checklist:
    body.append(f"<tr><td>{esc(step)}</td><td>{esc(title)}</td><td>{esc(detail)}</td></tr>")
body.append("</table><h2>Results by functionality</h2><table><tr><th>Id</th><th>Pass</th><th>Fail</th><th>Untested</th></tr>")
for func, rows in by_func.items():
    body.append("<tr><td>{}</td><td>{}</td><td>{}</td><td>{}</td></tr>".format(
        esc(func),
        sum(1 for row in rows if row["status"] == "PASS"),
        sum(1 for row in rows if row["status"] == "FAIL"),
        sum(1 for row in rows if row["status"] == "UNTESTED"),
    ))
body.append("</table><h2>Cases</h2>")
for row in results:
    klass = row["status"].lower()
    body.append(f"<h3>{esc(row['id'])} {esc(row['title'])} <span class='{klass}'>{esc(row['status'])}</span></h3>")
    body.append(f"<p><strong>Given</strong> {esc(row['given'])}<br><strong>When</strong> {esc(row['when'])}<br><strong>Then</strong> {esc(row['then'])}</p>")
    if row.get("detail"):
        body.append(f"<p>{esc(row['detail'])}</p>")
body.append("</body></html>")
(root / "test-plan.html").write_text("".join(body))
print("wrote", summary)
